"""
Phase 5D: exports a generated Paper to a .docx file with python-docx, in
three flavours:
  - "student": a clean exam copy -- header, general instructions, sections
    with their own instructions, numbered questions (either_or as "(a) ...
    OR (b) ..."), MCQ options lettered A-D, marks right-aligned. No
    Bloom's level, no document tags, no answers.
  - "teacher": the student copy plus a small grey tag line under each
    question ("[10 marks - Evaluate - <document>]") and a final page with
    the blueprint (target/planned/actual % per Bloom level, marks per
    document).
  - "answer_key": each question's number, correct answer (MCQ: letter +
    option text), and marking scheme points with their marks.

Deliberately takes the exact dict shape main.py's _paper_detail() already
returns (never touches the database itself) -- so this is trivially unit
-testable against a hand-built mock paper, and can never drift from what
the API/frontend already see. An empty slot (question is None) always
renders as "[Question missing]" here, in every kind -- it's the caller's
job (the export route) to decide whether that should block the request
instead (409 for a student copy with gaps).
"""
import io
import os

import bloom
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

EXPORT_KINDS = ("student", "teacher", "answer_key")

GENERAL_INSTRUCTIONS = (
    "Attempt all questions unless a section says otherwise. Write your answers "
    "in the space provided. Marks for each question are shown in brackets on "
    "the right."
)

MCQ_LETTERS = ("A", "B", "C", "D")
TRUE_FALSE_TEXT = {"A": "True", "B": "False"}

TAG_FONT_SIZE = Pt(8)
TAG_FONT_COLOR = RGBColor(0x80, 0x80, 0x80)


def build_docx(paper: dict, kind: str) -> io.BytesIO:
    """Returns an in-memory .docx (seeked to 0) for `paper` (a _paper_detail()-shaped dict). Raises ValueError for an unknown `kind`."""
    if kind not in EXPORT_KINDS:
        raise ValueError(f"Unknown export kind '{kind}'")

    doc = Document()
    _set_base_style(doc)
    _add_header(doc, paper)

    if kind == "answer_key":
        _add_answer_key(doc, paper)
    else:
        _add_general_instructions(doc)
        _add_sections(doc, paper, kind)

    if kind == "teacher":
        _add_blueprint_page(doc, paper)

    _add_page_number_footer(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# -----------------------------
# Styling
# -----------------------------
def _set_base_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def _add_page_number_footer(doc: Document) -> None:
    """
    "Page X of Y" via Word field codes (PAGE / NUMPAGES) -- python-docx
    can't compute the actual numbers itself, only embed the field codes;
    Word fills them in once the file is opened.
    """
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Page ")
    _add_field(paragraph, "PAGE")
    paragraph.add_run(" of ")
    _add_field(paragraph, "NUMPAGES")


def _add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _right_tab_stop(doc: Document):
    section = doc.sections[0]
    return section.page_width - section.left_margin - section.right_margin


# -----------------------------
# Shared data helpers
# -----------------------------
def _marks_on_paper(paper: dict) -> int:
    """Every distinct numbered question once each, regardless of choice -- matches paper.marks_on_paper()."""
    return sum(s["marks"] for s in paper.get("slots", []) if s.get("option") in (None, "a"))


def _total_marks(paper: dict) -> int:
    """
    What a student answering optimally actually scores -- matches
    paper.total_marks(): identical to _marks_on_paper() except an "any_n"
    section only counts `attempt` of its num_questions (the student never
    answers the rest). This, not _marks_on_paper(), is the right "Max
    Marks" figure on a printed paper (Phase 5E EXTRA fix) -- e.g. 50 for
    mid_sem_50's "any 4 of 5" section, not 55. _marks_on_paper() stays the
    basis for the Bloom blueprint's percentages, matching how plan_paper()
    computed the targets in the first place.
    """
    total = 0
    for section in (paper.get("pattern") or {}).get("sections", []):
        if section.get("choice") == "any_n" and section.get("attempt"):
            total += section["attempt"] * section["marks_each"]
        else:
            total += section["num_questions"] * section["marks_each"]
    return total


def _format_duration(minutes) -> str:
    """120 -> "2 hours", 90 -> "1 hour 30 minutes", 60 -> "1 hour", 45 -> "45 minutes"."""
    if not minutes:
        return ""
    hours, mins = divmod(int(minutes), 60)
    parts = []
    if hours:
        parts.append(f"{hours} hour" + ("s" if hours != 1 else ""))
    if mins:
        parts.append(f"{mins} minute" + ("s" if mins != 1 else ""))
    return " ".join(parts) if parts else "0 minutes"


def _group_slots(slots: list) -> list:
    """Groups slots by (section, number) so an either_or pair's two options stay together, in plan order."""
    groups = []
    index_by_key = {}
    for slot in slots:
        key = (slot["section"], slot["number"])
        if key not in index_by_key:
            index_by_key[key] = len(groups)
            groups.append([])
        groups[index_by_key[key]].append(slot)
    return groups


def _level_label(level_key) -> str:
    try:
        return bloom.get_level(level_key).label
    except ValueError:
        return level_key or "Unknown level"


def _strip_extension(filename: str) -> str:
    return os.path.splitext(filename)[0] if filename else filename


def _document_name(slot: dict) -> str:
    """The document's name, without its file extension (Phase 5E polish) -- e.g. "CN Notes", not "CN Notes.docx"."""
    question = slot.get("question")
    if question and question.get("source") and question["source"].get("book_name"):
        return _strip_extension(question["source"]["book_name"])
    if slot.get("book_id") is not None:
        return f"Document #{slot['book_id']}"
    return "Unknown document"


def _document_names_by_id(paper: dict) -> dict:
    """Best-effort {book_id (str): name} built from whichever filled slots know their document's real name."""
    names = {}
    for slot in paper.get("slots", []):
        if slot.get("book_id") is None:
            continue
        question = slot.get("question")
        if question and question.get("source") and question["source"].get("book_name"):
            names[str(slot["book_id"])] = _document_name(slot)
    return names


def _fmt_marks(value) -> str:
    value = value or 0
    return str(int(value)) if float(value) == int(value) else f"{value:.1f}"


# -----------------------------
# Header / instructions
# -----------------------------
def _add_header(doc: Document, paper: dict) -> None:
    title = doc.add_heading(paper.get("title") or "Untitled Paper", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if paper.get("subject"):
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.add_run(paper["subject"]).italic = True

    meta_parts = [f"Max Marks: {_total_marks(paper)}"]
    duration = (paper.get("pattern") or {}).get("duration_minutes")
    if duration:
        meta_parts.append(f"Duration: {_format_duration(duration)}")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("      ".join(meta_parts)).bold = True

    name_line = doc.add_paragraph()
    name_line.add_run("Name: ____________________________        Roll No.: ______________")

    doc.add_paragraph()


def _add_general_instructions(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run("Instructions: ").bold = True
    p.add_run(GENERAL_INSTRUCTIONS)
    doc.add_paragraph()


# -----------------------------
# Sections / questions (student + teacher)
# -----------------------------
def _add_sections(doc: Document, paper: dict, kind: str) -> None:
    slots_by_section = {}
    for slot in paper.get("slots", []):
        slots_by_section.setdefault(slot["section"], []).append(slot)

    for section_meta in (paper.get("pattern") or {}).get("sections", []):
        section_slots = slots_by_section.get(section_meta["key"])
        if not section_slots:
            continue

        doc.add_heading(section_meta.get("title") or f"Section {section_meta['key']}", level=1)
        instructions = section_meta.get("instructions")
        if instructions:
            doc.add_paragraph().add_run(instructions).italic = True

        for group in _group_slots(section_slots):
            _add_question_group(doc, group, kind)


def _add_question_group(doc: Document, group: list, kind: str) -> None:
    number = group[0]["number"]
    if len(group) == 1:
        _add_single_question(doc, group[0], f"{number}.", kind, indent=False)
        return

    heading = doc.add_paragraph()
    heading.add_run(f"{number}. Answer EITHER (a) OR (b):").bold = True
    for slot in group:
        _add_single_question(doc, slot, f"({slot['option']})", kind, indent=True)


def _add_single_question(doc: Document, slot: dict, label: str, kind: str, indent: bool) -> None:
    question = slot.get("question")

    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.8)
    p.add_run(f"{label} ").bold = True

    if question is None:
        p.add_run("[Question missing]").italic = True
    else:
        p.add_run(question["question"])

    p.paragraph_format.tab_stops.add_tab_stop(_right_tab_stop(doc), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(f"\t[{slot['marks']}]")

    if question and question.get("options"):
        option_indent = Cm(1.2) if indent else Cm(0.6)
        for idx, option_text in enumerate(question["options"]):
            letter = MCQ_LETTERS[idx] if idx < len(MCQ_LETTERS) else str(idx + 1)
            opt_p = doc.add_paragraph(f"{letter}. {option_text}")
            opt_p.paragraph_format.left_indent = option_indent

    if kind == "teacher" and question is not None:
        tag_p = doc.add_paragraph()
        if indent:
            tag_p.paragraph_format.left_indent = Cm(0.8)
        tag_run = tag_p.add_run(
            f"[{slot['marks']} marks · {_level_label(question['bloom_level'])} · {_document_name(slot)}]"
        )
        tag_run.font.size = TAG_FONT_SIZE
        tag_run.font.color.rgb = TAG_FONT_COLOR
        tag_run.italic = True

    doc.add_paragraph()


# -----------------------------
# Answer key
# -----------------------------
def _add_answer_key(doc: Document, paper: dict) -> None:
    doc.add_heading("Answer Key", level=1)

    for group in _group_slots(paper.get("slots", [])):
        for slot in group:
            option_suffix = f" ({slot['option']})" if slot.get("option") else ""
            question = slot.get("question")

            p = doc.add_paragraph()
            p.add_run(f"{slot['number']}{option_suffix}. ").bold = True

            if question is None:
                p.add_run("[Question missing]").italic = True
                doc.add_paragraph()
                continue

            p.add_run(_answer_text(question))

            for point in question.get("marking_scheme") or []:
                bullet = doc.add_paragraph(style="List Bullet")
                bullet.add_run(f"{point['point']} ({point['marks']})")

            doc.add_paragraph()


def _answer_text(question: dict) -> str:
    correct = question.get("correct_answer") or ""
    if question["type"] == "mcq" and question.get("options"):
        idx = MCQ_LETTERS.index(correct) if correct in MCQ_LETTERS else None
        option_text = question["options"][idx] if idx is not None and idx < len(question["options"]) else ""
        return f"{correct}. {option_text}".strip()
    if question["type"] == "true_false":
        return TRUE_FALSE_TEXT.get(correct, correct)
    return correct


# -----------------------------
# Blueprint page (teacher only)
# -----------------------------
def _pct(marks, total) -> str:
    if not total:
        return "0%"
    return f"{round((marks or 0) / total * 100)}%"


def _fill_blueprint_cell(cell, marks, total) -> None:
    """Main number is the %-of-paper (what item 2 of the Phase 5 fixes asked to lead with); the exact
    marks, which is what a fixed-size question bank can actually hit, sit underneath in small grey text."""
    cell.text = ""
    cell.paragraphs[0].add_run(_pct(marks, total)).bold = True
    sub_run = cell.add_paragraph().add_run(f"{_fmt_marks(marks)} of {_fmt_marks(total)} printed marks")
    sub_run.font.size = TAG_FONT_SIZE
    sub_run.font.color.rgb = TAG_FONT_COLOR


def _add_blueprint_page(doc: Document, paper: dict) -> None:
    doc.add_page_break()
    doc.add_heading("Blueprint", level=1)

    # Marks on paper (every printed question once, incl. unattempted any_n
    # choices) differs from the "Max Marks" on the header (what a student
    # actually answers) whenever any section offers a choice -- labelled
    # here so that distinction stays visible to the teacher.
    marks_on_paper = _marks_on_paper(paper)
    total_marks = _total_marks(paper)
    if marks_on_paper != total_marks:
        doc.add_paragraph().add_run(
            f"Marks printed on paper: {marks_on_paper} (students answer {total_marks})"
        ).italic = True

    doc.add_heading("Marks by Bloom's level", level=2)
    target = paper.get("blueprint_target") or {}
    planned = (paper.get("blueprint_planned") or {}).get("by_level") or {}
    actual = (paper.get("blueprint_actual") or {}).get("by_level") or {}
    level_total = sum(target.values()) or marks_on_paper or 1

    level_table = doc.add_table(rows=1, cols=4)
    level_table.style = "Light Grid Accent 1"
    header_cells = level_table.rows[0].cells
    header_cells[0].text, header_cells[1].text, header_cells[2].text, header_cells[3].text = (
        "Level", "Target", "Planned", "Actual"
    )
    for level_key in bloom.level_keys():
        row = level_table.add_row().cells
        row[0].text = _level_label(level_key)
        _fill_blueprint_cell(row[1], target.get(level_key), level_total)
        _fill_blueprint_cell(row[2], planned.get(level_key), level_total)
        _fill_blueprint_cell(row[3], actual.get(level_key), level_total)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Exact match isn't always possible because each question is worth fixed marks (2, 5 or 10)."
    )
    note_run.italic = True
    note_run.font.size = TAG_FONT_SIZE
    note_run.font.color.rgb = TAG_FONT_COLOR

    doc.add_paragraph()
    doc.add_heading("Marks by document", level=2)
    by_unit = (paper.get("blueprint_actual") or {}).get("by_unit") or {}
    names = _document_names_by_id(paper)

    unit_table = doc.add_table(rows=1, cols=2)
    unit_table.style = "Light Grid Accent 1"
    unit_header = unit_table.rows[0].cells
    unit_header[0].text, unit_header[1].text = "Document", "Marks"
    for unit_id, marks in by_unit.items():
        row = unit_table.add_row().cells
        row[0].text = names.get(str(unit_id), f"Document #{unit_id}")
        row[1].text = _fmt_marks(marks)

    warnings = paper.get("warnings") or []
    if warnings:
        doc.add_paragraph()
        doc.add_heading("Warnings", level=2)
        for warning in warnings:
            doc.add_paragraph(warning, style="List Bullet")
