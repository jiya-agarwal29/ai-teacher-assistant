"""
Tests for paper_export.py's build_docx() -- pure, no DB/app needed. Each
test hands it a hand-built dict shaped exactly like main.py's
_paper_detail() return value, then reads the generated .docx back with
python-docx to check its text.
"""
import pytest
from docx import Document

import paper_export


def _question(
    question_id, q_type, level, text, correct_answer, marks, marking_scheme,
    options=None, book_name="CN Notes.docx", page_number=1, bloom_check="verified", detected_level=None
):
    return {
        "question_id": question_id, "topic": "networking", "type": q_type,
        "bloom_level": level, "bloom_label": level.title(), "bloom_check": bloom_check,
        "detected_level": detected_level, "question": text, "options": options,
        "correct_answer": correct_answer, "explanation": "An explanation.",
        "marks": marks, "marking_scheme": marking_scheme,
        "created_at": "2026-01-01T00:00:00",
        "source": {"book_name": book_name, "page_number": page_number},
    }


def _mock_paper():
    return {
        "title": "Mid-Sem Mock Paper",
        "subject": "Networks",
        "status": "ready",
        "pattern": {
            "duration_minutes": 120,
            "sections": [
                {
                    "key": "A", "title": "Section A", "instructions": "Answer all questions.",
                    "marks_each": 2, "num_questions": 2, "choice": "none", "attempt": None
                },
                {
                    "key": "B", "title": "Section B", "instructions": "Answer EITHER (a) OR (b) for each question.",
                    "marks_each": 10, "num_questions": 1, "choice": "either_or", "attempt": None
                },
            ],
        },
        "blueprint_target": {
            "remember": 2.0, "understand": 2.0, "apply": 0, "analyze": 10.0, "evaluate": 0, "create": 0
        },
        "blueprint_actual": {
            "by_level": {"remember": 2, "understand": 0, "apply": 0, "analyze": 10, "evaluate": 0, "create": 0},
            "by_unit": {"1": 2, "2": 10},
        },
        "warnings": ["Section 'B' question 3 (b): classified as Evaluate, not the intended Analyze."],
        "slots": [
            {
                "id": 1, "section": "A", "number": "1", "option": None, "marks": 2,
                "bloom_level": "remember", "book_id": 1,
                "question": _question(
                    101, "mcq", "remember", "What does CRC stand for?", "A", 2,
                    [{"point": "Identifies Cyclic Redundancy Check", "marks": 2}],
                    options=["Cyclic Redundancy Check", "Code Redundancy Check", "Cyclic Record Check", "Cyclic Reliability Check"],
                ),
            },
            {
                "id": 2, "section": "A", "number": "2", "option": None, "marks": 2,
                "bloom_level": "understand", "book_id": 1, "question": None
            },
            {
                "id": 3, "section": "B", "number": "3", "option": "a", "marks": 10,
                "bloom_level": "analyze", "book_id": 2,
                "question": _question(
                    102, "long_answer", "analyze", "Compare TCP and UDP.", "TCP is reliable; UDP is not.", 10,
                    [{"point": "TCP reliability", "marks": 5}, {"point": "UDP simplicity", "marks": 5}],
                    book_name="NS Notes.pdf", page_number=2,
                ),
            },
            {
                "id": 4, "section": "B", "number": "3", "option": "b", "marks": 10,
                "bloom_level": "analyze", "book_id": 2,
                "question": _question(
                    103, "scenario", "analyze", "Examine firewall placement in the topology.",
                    "Place the firewall after the edge router.", 10,
                    [{"point": "Correct placement", "marks": 10}],
                    book_name="NS Notes.pdf", page_number=3,
                    bloom_check="flagged", detected_level="evaluate",
                ),
            },
        ],
    }


def _docx_text(buffer):
    doc = Document(buffer)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


@pytest.mark.parametrize("kind", paper_export.EXPORT_KINDS)
def test_build_docx_does_not_raise(kind):
    buffer = paper_export.build_docx(_mock_paper(), kind)
    assert buffer.getbuffer().nbytes > 0


def test_unknown_kind_raises_value_error():
    with pytest.raises(ValueError):
        paper_export.build_docx(_mock_paper(), "bogus")


def test_student_copy_has_no_level_names_document_tags_or_answers():
    text = _docx_text(paper_export.build_docx(_mock_paper(), "student"))
    assert "Remember" not in text
    assert "Analyze" not in text
    assert "marking scheme" not in text.lower()
    assert "CN Notes.docx" not in text
    assert "NS Notes.pdf" not in text
    # Free-text correct answers must never appear in the student copy.
    assert "TCP is reliable; UDP is not." not in text
    assert "Place the firewall after the edge router." not in text
    # Header basics are present.
    assert "Mid-Sem Mock Paper" in text
    assert "Networks" in text
    assert "Max Marks: 14" in text
    assert "Duration: 2 hours" in text
    assert "Roll No." in text


def test_student_copy_shows_mcq_options_lettered():
    text = _docx_text(paper_export.build_docx(_mock_paper(), "student"))
    assert "A. Cyclic Redundancy Check" in text
    assert "B. Code Redundancy Check" in text
    assert "C. Cyclic Record Check" in text
    assert "D. Cyclic Reliability Check" in text


def test_either_or_shows_both_options_with_a_b_labels():
    text = _docx_text(paper_export.build_docx(_mock_paper(), "student"))
    assert "Answer EITHER (a) OR (b)" in text
    assert "(a)" in text
    assert "(b)" in text
    assert "Compare TCP and UDP." in text
    assert "Examine firewall placement in the topology." in text


def test_teacher_copy_has_tags_and_blueprint():
    text = _docx_text(paper_export.build_docx(_mock_paper(), "teacher"))
    assert "Remember" in text
    assert "Analyze" in text
    assert "CN Notes" in text
    assert "NS Notes" in text
    assert "Blueprint" in text
    assert "Marks by Bloom's level" in text
    assert "Marks by document" in text
    assert "classified as Evaluate" in text  # paper-level warning included


def test_empty_slot_shows_question_missing_in_teacher_and_answer_key():
    for kind in ("teacher", "answer_key"):
        text = _docx_text(paper_export.build_docx(_mock_paper(), kind))
        assert "[Question missing]" in text


def test_student_copy_also_shows_question_missing_placeholder():
    # build_docx itself never blocks on a gap -- only the export route
    # decides whether a gap should 409 instead of calling this at all.
    text = _docx_text(paper_export.build_docx(_mock_paper(), "student"))
    assert "[Question missing]" in text


def test_answer_key_has_every_question_and_scheme_point():
    text = _docx_text(paper_export.build_docx(_mock_paper(), "answer_key"))
    assert "Answer Key" in text
    assert "A. Cyclic Redundancy Check" in text
    assert "Identifies Cyclic Redundancy Check (2)" in text
    assert "TCP reliability (5)" in text
    assert "UDP simplicity (5)" in text
    assert "Correct placement (10)" in text


def test_true_false_answer_key_shows_true_or_false_word():
    paper = _mock_paper()
    paper["slots"][0]["question"]["type"] = "true_false"
    paper["slots"][0]["question"]["options"] = ["True", "False"]
    paper["slots"][0]["question"]["correct_answer"] = "B"
    text = _docx_text(paper_export.build_docx(paper, "answer_key"))
    assert "False" in text


# -----------------------------
# Phase 5E EXTRA: total_marks (not marks_on_paper), duration wording,
# document names without their file extension
# -----------------------------
def test_header_shows_total_marks_not_marks_on_paper_for_any_n_sections():
    # mid_sem_50-style "attempt any 2 of 3" section: a student only ever
    # answers 2 * 5 = 10 marks, not the 15 marks printed on the page.
    paper = {
        "title": "Any-N Test Paper", "subject": None, "status": "ready",
        "pattern": {
            "duration_minutes": 60,
            "sections": [
                {
                    "key": "A", "title": "Section A", "instructions": "Attempt any 2 of the following 3 questions.",
                    "marks_each": 5, "num_questions": 3, "choice": "any_n", "attempt": 2
                },
            ],
        },
        "blueprint_target": {}, "blueprint_actual": {"by_level": {}, "by_unit": {}}, "warnings": [],
        "slots": [
            {
                "id": n, "section": "A", "number": str(n), "option": None, "marks": 5,
                "bloom_level": "remember", "book_id": 1,
                "question": _question(n, "short_answer", "remember", f"Question {n} text", f"Answer {n}", 5,
                                       [{"point": "x", "marks": 5}])
            }
            for n in (1, 2, 3)
        ],
    }
    text = _docx_text(paper_export.build_docx(paper, "student"))
    assert "Max Marks: 10" in text  # attempt(2) * marks_each(5), not 3 * 5 = 15

    # The teacher copy's blueprint page may still show the 15-mark
    # "marks on paper" figure, explicitly labelled as such.
    teacher_text = _docx_text(paper_export.build_docx(paper, "teacher"))
    assert "Marks on paper (all printed questions): 15" in teacher_text


@pytest.mark.parametrize("minutes,expected", [(120, "2 hours"), (90, "1 hour 30 minutes"), (60, "1 hour"), (45, "45 minutes")])
def test_duration_is_worded_not_raw_minutes(minutes, expected):
    paper = _mock_paper()
    paper["pattern"]["duration_minutes"] = minutes
    text = _docx_text(paper_export.build_docx(paper, "student"))
    assert f"Duration: {expected}" in text


def test_teacher_tags_show_document_name_without_extension():
    text = _docx_text(paper_export.build_docx(_mock_paper(), "teacher"))
    assert "CN Notes ·" in text or "· CN Notes]" in text
    assert "CN Notes.docx" not in text
    assert "NS Notes.pdf" not in text
